"""Tests for invoice file parsing (PDF/DOCX upload -> extracted invoice fields).

The LLM extraction step is patched out so these run offline; they cover file
validation, text extraction, the heuristic fallback and the /parse endpoint.
"""
import io
from unittest.mock import patch

import pytest
from docx import Document

from src import invoice_parser
from src.invoice_parser import (
    extract_text_from_invoice,
    parse_invoice_text,
    validate_invoice_upload,
)

INVOICE_LINES = [
    "Lena Park Illustration",
    "Invoice #INV-042",
    "Bill To: Client Co.",
    "sam@clientco.com",
    "Due Date: 2026-04-02",
    "Rush delivery artwork",
    "Balance Due: $5,000.00",
]


def make_text_pdf(lines: list[str]) -> bytes:
    """Build a minimal one-page PDF whose text pypdf can extract."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    ops = ["BT", "/F1 12 Tf", "72 720 Td"]
    for i, line in enumerate(lines):
        if i:
            ops.append("0 -18 Td")
        ops.append(f"({esc(line)}) Tj")
    ops.append("ET")
    return _assemble_pdf("\n".join(ops).encode("latin-1"))


def make_image_only_pdf() -> bytes:
    """A PDF page that draws a shape but contains no text, like a scan."""
    return _assemble_pdf(b"0 0 1 rg 72 72 200 200 re f")


def _assemble_pdf(stream: bytes) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for num, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{num} 0 obj\n".encode() + body + b"\nendobj\n")
    xref_at = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str], table_rows: list[list[str]]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
    for r, row in enumerate(table_rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def llm_unavailable(_text):
    raise RuntimeError("connection error")


# ─── File validation ─────────────────────────────────────────────────────────

@pytest.mark.parametrize("filename,content_type,expected", [
    ("invoice.pdf", "application/pdf", "pdf"),
    ("INVOICE.PDF", None, "pdf"),
    ("invoice.docx", None, "docx"),
])
def test_accepts_pdf_and_docx(filename, content_type, expected):
    assert validate_invoice_upload(filename, content_type) == expected


@pytest.mark.parametrize("filename,message", [
    ("invoice.doc", "Legacy .doc"),
    ("scan.png", "Unsupported invoice file"),
])
def test_rejects_other_files_with_a_clear_message(filename, message):
    with pytest.raises(ValueError, match=message):
        validate_invoice_upload(filename, None)


# ─── Text extraction ─────────────────────────────────────────────────────────

def test_extracts_text_from_a_text_pdf():
    text = extract_text_from_invoice(make_text_pdf(INVOICE_LINES), "pdf")
    assert "INV-042" in text
    assert "sam@clientco.com" in text
    assert "Balance Due: $5,000.00" in text


def test_image_only_pdf_explains_scans_are_not_supported():
    with pytest.raises(ValueError, match="Scanned/image-only PDFs are not supported"):
        extract_text_from_invoice(make_image_only_pdf(), "pdf")


def test_extracts_paragraphs_and_tables_from_docx():
    content = make_docx(
        ["Invoice #INV-077", "Bill To: Studio North"],
        [["Item", "Amount"], ["Design retainer", "$2,400.00"]],
    )
    text = extract_text_from_invoice(content, "docx")
    assert "INV-077" in text
    assert "Design retainer | $2,400.00" in text


# ─── Field extraction ────────────────────────────────────────────────────────

def test_falls_back_to_heuristics_when_the_llm_is_unavailable():
    text = extract_text_from_invoice(make_text_pdf(INVOICE_LINES), "pdf")
    with patch.object(invoice_parser, "_llm_parse", side_effect=llm_unavailable):
        parsed = parse_invoice_text(text)

    assert parsed.invoice_id == "INV-042"
    assert parsed.client_name == "Client Co."
    assert parsed.client_email == "sam@clientco.com"
    assert parsed.invoice_amount == 5000.0
    assert parsed.due_date == "2026-04-02"
    assert parsed.days_overdue is not None and parsed.days_overdue >= 0
    assert "AI extraction failed; basic document parsing was used instead." in parsed.warnings


def test_fallback_prefers_the_bill_to_email_over_the_senders():
    # Invoices usually print the sender's own email in the header, above Bill To.
    lines = ["Lena Park Illustration", "lena@lenapark.studio", *INVOICE_LINES[1:]]
    text = extract_text_from_invoice(make_text_pdf(lines), "pdf")
    with patch.object(invoice_parser, "_llm_parse", side_effect=llm_unavailable):
        parsed = parse_invoice_text(text)

    assert parsed.client_email == "sam@clientco.com"


def test_llm_result_wins_and_heuristics_fill_gaps():
    text = extract_text_from_invoice(make_text_pdf(INVOICE_LINES), "pdf")
    llm_result = invoice_parser.ParsedInvoice(
        invoice_id="INV-042", client_name="Client Co.", invoice_amount=5000.0,
        jurisdiction="Oregon", confidence=0.9,
    )
    with patch.object(invoice_parser, "_llm_parse", return_value=llm_result):
        parsed = parse_invoice_text(text)

    assert parsed.jurisdiction == "Oregon"
    assert parsed.confidence == 0.9
    # Missing from the LLM answer, recovered from the document text.
    assert parsed.client_email == "sam@clientco.com"
    assert parsed.due_date == "2026-04-02"


# ─── /api/invoices/parse endpoint ────────────────────────────────────────────

@pytest.mark.asyncio
async def test_parse_endpoint_returns_fields_for_an_uploaded_pdf(client):
    with patch.object(invoice_parser, "_llm_parse", side_effect=llm_unavailable):
        r = await client.post(
            "/api/invoices/parse",
            files={"file": ("INV-042.pdf", make_text_pdf(INVOICE_LINES), "application/pdf")},
        )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["extracted"]["invoice_id"] == "INV-042"
    assert body["extracted"]["client_email"] == "sam@clientco.com"
    assert body["extracted"]["invoice_amount"] == 5000.0
    assert body["missing_fields"] == []
    assert body["file"]["name"] == "INV-042.pdf"


@pytest.mark.asyncio
async def test_parse_endpoint_rejects_a_scanned_pdf_with_400(client):
    r = await client.post(
        "/api/invoices/parse",
        files={"file": ("scan.pdf", make_image_only_pdf(), "application/pdf")},
    )
    assert r.status_code == 400
    assert "Scanned/image-only PDFs are not supported" in r.json()["detail"]


@pytest.mark.asyncio
async def test_parse_endpoint_rejects_unsupported_files_with_400(client):
    r = await client.post(
        "/api/invoices/parse",
        files={"file": ("photo.png", b"\x89PNG fake", "image/png")},
    )
    assert r.status_code == 400
    assert "Unsupported invoice file" in r.json()["detail"]


@pytest.mark.asyncio
@pytest.mark.parametrize("name,content,mime", [
    ("fake.pdf", b"not a pdf at all", "application/pdf"),
    ("fake.docx", b"not a zip either", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
])
async def test_parse_endpoint_rejects_corrupt_files_with_400(client, name, content, mime):
    """A renamed or corrupt file is a user error (400 with a clear message), not a server error."""
    r = await client.post("/api/invoices/parse", files={"file": (name, content, mime)})
    assert r.status_code == 400, r.text
    assert "readable" in r.json()["detail"].lower()
